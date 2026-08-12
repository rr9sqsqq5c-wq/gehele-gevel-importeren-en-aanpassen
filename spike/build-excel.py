# Bouwt het Excel-werkboek uit excel-data.json (per tabblad een dataset).
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
data = json.load(open(os.path.join(OUT, 'excel-data.json'), encoding='utf-8'))
XLSX = os.path.join(OUT, 'brickboard-extractie.xlsx')

HEAD = Font(bold=True, color='FFFFFF')
HFILL = PatternFill('solid', fgColor='2F5496')
TITLE = Font(bold=True, size=13)
SUB = Font(italic=True, size=9, color='808080')
thin = Side(style='thin', color='D9D9D9')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
NL = {'RED': 'Rood', 'GREEN': 'Groen', 'PENNANTS': 'Penanten'}

wb = Workbook()

def sheet(name):
    ws = wb.create_sheet(name)
    return ws

def table(ws, headers, rows, start_row=1, widths=None, numfmt=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(start_row, c, h); cell.font = HEAD; cell.fill = HFILL; cell.alignment = Alignment(horizontal='center'); cell.border = BORDER
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v); cell.border = BORDER
            if numfmt and c in numfmt: cell.number_format = numfmt[c]
    if widths:
        for c, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    return ws

# ── Samenvatting ──────────────────────────────────────────────────────────────
ws = wb.active; ws.title = 'Samenvatting'
ws['A1'] = 'Brickboard — extractie uit 3 IFC (RED / GREEN / PENNANTS)'; ws['A1'].font = TITLE
ws['A2'] = 'Panelen = beklede gevelpanelen (verticale platen; 1 plaat = 1 paneel). Maten in mm, oppervlak in m².'; ws['A2'].font = SUB
rows = [[NL.get(s['deel'], s['deel']), s['panelen'], s['paneel_m2'], s['steenstrips'], s['strip_m2'], s['latten'], s['latten_m'], s['gevels']] for s in data['samenvatting']]
tot = ['TOTAAL', sum(s['panelen'] for s in data['samenvatting']), round(sum(s['paneel_m2'] for s in data['samenvatting']), 1),
       sum(s['steenstrips'] for s in data['samenvatting']), round(sum(s['strip_m2'] for s in data['samenvatting']), 1),
       sum(s['latten'] for s in data['samenvatting']), round(sum(s['latten_m'] for s in data['samenvatting']), 1), '']
rows.append(tot)
table(ws, ['Gebouwdeel', 'Panelen', 'Paneel m²', 'Steenstrips', 'Strip m²', 'Latten', 'Latten m', 'Gevels'], rows, start_row=4,
      widths=[16, 12, 11, 13, 11, 10, 11, 9])
for c in range(1, 9):
    ws.cell(4 + len(rows), c).font = Font(bold=True)

# ── Bestellen (hele strippen) + zaagverlies m² ────────────────────────────────
ws = sheet('Bestellen strippen')
ws['A1'] = 'Te bestellen: hele steenstrips (221 mm) — met zaagverlies in m²'; ws['A1'].font = TITLE
ws['A2'] = 'Best-fit zaagoptimalisatie, zaagsnede 3 mm. Rood en Groen aparte kleuren (penanten zijn groen). Oppervlak = lengte × striphoogte 51 mm.'; ws['A2'].font = SUB
tb = data['bestellen']
rows = [[b['kleur'], b['hele_strippen'], b['ingekocht_m2'], b['gevel_m2'], b['verlies_m2'], b['verlies_pct'] / 100, b['zaagsnede_m2'], b['rest_m2']] for b in tb]
S = lambda k: round(sum(b[k] for b in tb), 1)
rows.append(['TOTAAL', sum(b['hele_strippen'] for b in tb), S('ingekocht_m2'), S('gevel_m2'), S('verlies_m2'),
             (sum(b['verlies_m2'] for b in tb) / sum(b['ingekocht_m2'] for b in tb)), S('zaagsnede_m2'), S('rest_m2')])
table(ws, ['Kleur', 'Hele strippen', 'Ingekocht (m²)', 'Op de gevel (m²)', 'Zaagverlies (m²)', 'Zaagverlies %', 'waarvan zaagsnede (m²)', 'waarvan rest (m²)'],
      rows, start_row=4, widths=[24, 14, 14, 16, 16, 13, 21, 17], numfmt={6: '0.0%'})
for c in range(1, 9):
    ws.cell(4 + len(rows), c).font = Font(bold=True)

# ── Panelen RED / GREEN ───────────────────────────────────────────────────────
for key, label in [('panelenRED', 'Panelen Rood'), ('panelenGREEN', 'Panelen Groen')]:
    ws = sheet(label)
    ws['A1'] = f'{label} — genummerd per gevel (onder→boven, links→rechts)'; ws['A1'].font = TITLE
    rows = [[p['nummer'], p['gevel'], p['breedte'], p['hoogte'], p['elementen']] for p in data[key]]
    table(ws, ['Paneelnummer', 'Gevel', 'Breedte (mm)', 'Hoogte (mm)', '# elementen'], rows, start_row=3, widths=[14, 8, 13, 13, 12])

# ── Penanten ──────────────────────────────────────────────────────────────────
ws = sheet('Penanten')
ws['A1'] = 'Penanten (verticale kolommen) — genummerd per gevel'; ws['A1'].font = TITLE
rows = [[p['nummer'], p['gevel'], p['type'], p['zijden'], p['elementen'], p['x'], p['y']] for p in data['penanten']]
table(ws, ['Nummer', 'Gevel', 'Type', 'Zijden', '# elementen', 'X (mm)', 'Y (mm)'], rows, start_row=3, widths=[10, 8, 8, 8, 12, 10, 10])
r0 = 3 + len(rows) + 3
ws.cell(r0 - 1, 1, 'Unieke penant-types').font = TITLE
trows = [[t['type'], t['aantal'], t['zijden'], t['strips'], t['boards'], t['latten']] for t in data['penantTypes']]
table(ws, ['Type', 'Aantal', 'Zijden', 'Strips', 'Platen', 'Latten'], trows, start_row=r0)

# ── Steenstrips (maten) ───────────────────────────────────────────────────────
ws = sheet('Steenstrips maten')
ws['A1'] = 'Steenstrips per maat (lengte × hoogte × dikte)'; ws['A1'].font = TITLE
rows = []
for k in ('RED', 'GREEN', 'PENNANTS'):
    for s in data['steenstrips'][k]:
        rows.append([NL[k], s['lengte'], s['hoogte'], s['dikte'], s['aantal']])
table(ws, ['Gebouwdeel', 'Lengte (mm)', 'Hoogte (mm)', 'Dikte (mm)', 'Aantal'], rows, start_row=3, widths=[14, 12, 12, 11, 10])

# ── Latten (maten) ────────────────────────────────────────────────────────────
ws = sheet('Latten maten')
ws['A1'] = 'Latten per maat (lengte × breedte × dikte)'; ws['A1'].font = TITLE
rows = []
for k in ('RED', 'GREEN', 'PENNANTS'):
    for s in data['latten'][k]:
        rows.append([NL[k], s['lengte'], s['breedte'], s['dikte'], s['aantal']])
table(ws, ['Gebouwdeel', 'Lengte (mm)', 'Breedte (mm)', 'Dikte (mm)', 'Aantal'], rows, start_row=3, widths=[14, 12, 12, 11, 10])

# ── Zaagschema Rood / Groen ───────────────────────────────────────────────────
for key, label in [('zaagRood', 'Zaagschema Rood'), ('zaagGroen', 'Zaagschema Groen')]:
    ws = sheet(label)
    ws['A1'] = f'{label} — welke stukken uit één 221 mm-strip'; ws['A1'].font = TITLE
    ws['A2'] = 'Rest = afval per strip (mm). Zaagsnede 3 mm.'; ws['A2'].font = SUB
    rows = [[p['patroon'], p['stukken'], p['strippen'], p['rest_mm']] for p in data[key]]
    table(ws, ['Zaagpatroon', 'Stukken/strip', 'Aantal strippen', 'Rest (mm)'], rows, start_row=4, widths=[30, 13, 15, 10])

wb.save(XLSX)
print('geschreven:', XLSX)
print('tabbladen:', wb.sheetnames)
