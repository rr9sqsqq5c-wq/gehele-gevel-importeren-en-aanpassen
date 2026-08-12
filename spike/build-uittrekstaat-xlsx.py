# Bouwt een Excel-uittrekstaat uit de CSV's van extract-green-building.mjs
import csv, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
XLSX = os.path.join(OUT, "green-building-uittrekstaat.xlsx")

def num(s):
    s = (s or "").strip().replace(",", ".")
    try:
        f = float(s)
        return int(f) if f == int(f) else f
    except ValueError:
        return s

HEAD = Font(bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor="2F5496")
TITLE = Font(bold=True, size=13)
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

CATS = [("panelen", "Panelen"), ("steenstrips", "Steenstrips"), ("latten", "Latten")]

wb = Workbook()

# ---- Samenvatting ----
ws = wb.active
ws.title = "Samenvatting"
ws["A1"] = "Uittrekstaat — 28072026-TEST-green-building-5.ifc"
ws["A1"].font = TITLE
ws["A2"] = "Elk item = rechthoekig profiel (lengte x breedte) geëxtrudeerd over de dikte. 'poly' = niet-rechthoekig, maat = bounding box."
ws["A2"].font = Font(italic=True, size=9, color="808080")
hdr = ["Categorie", "Aantal stuks", "Unieke maten", "Totaal vlak (m²)", "Totale lengte (m)"]
for c, h in enumerate(hdr, 1):
    cell = ws.cell(4, c, h); cell.font = HEAD; cell.fill = HEADFILL; cell.border = BORDER

def read_group(cat):
    rows = []
    with open(os.path.join(OUT, f"{cat}-uittrekstaat.csv"), encoding="utf-8") as f:
        r = csv.reader(f, delimiter=";"); next(r)
        for row in r:
            if not row: continue
            rows.append([num(row[0]), num(row[1]), num(row[2]), num(row[3]), row[4]])  # aantal,L,B,D,vorm
    return rows

def read_items(cat):
    rows = []
    with open(os.path.join(OUT, f"{cat}-per-stuk.csv"), encoding="utf-8") as f:
        r = csv.reader(f, delimiter=";"); head = next(r)
        for row in r:
            if not row: continue
            rows.append([num(row[0]), num(row[1]), num(row[2]), num(row[3]), num(row[4]), num(row[5]), num(row[6]), row[7]])
    return head, rows

srow = 5
for key, label in CATS:
    g = read_group(key)
    stuks = sum(x[0] for x in g)
    m2 = sum(x[0] * x[1] * x[2] for x in g) / 1e6
    lm = sum(x[0] * x[1] for x in g) / 1000
    vals = [label, stuks, len(g), round(m2, 1), round(lm, 1)]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(srow, c, v); cell.border = BORDER
        if c == 1: cell.font = Font(bold=True)
    srow += 1
for col, w in zip("ABCDE", (16, 14, 14, 16, 17)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"

# ---- per categorie: uittrekstaat + detail ----
def style_header(ws, headers, row=1):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row, c, h); cell.font = HEAD; cell.fill = HEADFILL
        cell.alignment = Alignment(horizontal="center"); cell.border = BORDER

for key, label in CATS:
    ws = wb.create_sheet(label)
    ws["A1"] = f"{label} — uittrekstaat (gegroepeerd op maat, meest voorkomend eerst)"
    ws["A1"].font = TITLE
    heads = ["Aantal", "Lengte (mm)", "Breedte (mm)", "Dikte (mm)", "Vorm"]
    style_header(ws, heads, row=3)
    r = 4
    for row in read_group(key):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v); cell.border = BORDER
            if c == 1: cell.font = Font(bold=True)
        r += 1
    for col, w in zip("ABCDE", (10, 13, 13, 11, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    # detail-tabblad
    wsd = wb.create_sheet(f"{label}-detail")
    head, items = read_items(key)
    nlheads = ["Solid-ID", "Product-ID", "Lengte (mm)", "Breedte (mm)", "Dikte (mm)", "X (mm)", "Y (mm)", "Vorm"]
    style_header(wsd, nlheads, row=1)
    for i, row in enumerate(items, 2):
        for c, v in enumerate(row, 1):
            wsd.cell(i, c, v)
    for col, w in zip("ABCDEFGH", (10, 11, 12, 12, 10, 10, 10, 8)):
        wsd.column_dimensions[col].width = w
    wsd.freeze_panes = "A2"

wb.save(XLSX)
print("geschreven:", XLSX)
print("tabbladen:", wb.sheetnames)
